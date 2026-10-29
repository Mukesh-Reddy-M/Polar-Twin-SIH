from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = 'docs/PolarTwin_Project_Building_Plan.docx'
NAVY = '102A43'
BLUE = '147D9B'
PALE = 'EAF4F7'
SLATE = '526477'
WHITE = 'FFFFFF'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=NAVY, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    header_pr = t.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement('w:tblHeader')
    repeat_header.set(qn('w:val'), 'true')
    header_pr.append(repeat_header)
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, WHITE, 8.5)
        shade(t.rows[0].cells[i], NAVY)
    for row_idx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, False, NAVY, 8.5)
            if row_idx % 2 == 0:
                shade(cells[i], 'F2F7FA')
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def numbered(doc, title, detail):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title + ' — ')
    r.bold = True
    p.add_run(detail)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.68)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.78)
sec.right_margin = Inches(.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(9)
styles['Normal'].font.color.rgb = RGBColor.from_string('263746')
styles['Normal'].paragraph_format.space_after = Pt(5)
for name, size, color in [('Title', 27, NAVY), ('Heading 1', 17, NAVY), ('Heading 2', 12, BLUE), ('Heading 3', 10, NAVY)]:
    s = styles[name]
    s.font.name = 'Aptos Display' if name in ('Title', 'Heading 1') else 'Aptos'
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
    s.paragraph_format.space_after = Pt(5)

# Header and footer
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run('POLAR TWIN  /  PROJECT BUILDING PLAN')
r.font.name = 'Aptos'; r.font.size = Pt(7.5); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('SIH 2026  •  Problem Statement 26060  •  Prototype roadmap')
r.font.name = 'Aptos'; r.font.size = Pt(7.5); r.font.color.rgb = RGBColor.from_string(SLATE)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.paragraph_format.space_after = Pt(4)
r = p.add_run('POLAR TWIN')
r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(BLUE); r.font.letter_spacing = Pt(1)
doc.add_paragraph('Project Building Plan', style='Title')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(13)
r = p.add_run('A phased plan to turn the Antarctic station prototype into a reliable, explainable operations demo')
r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string(SLATE)

callout = doc.add_table(rows=1, cols=1)
callout.style = 'Table Grid'
cell = callout.cell(0, 0); shade(cell, PALE)
set_cell_text(cell, 'PROJECT AIM  |  Demonstrate how a storm or equipment fault changes station conditions, triggers an explainable alert, creates a fuel requisition, and moves through approval, dispatch, and delivery while remaining useful during a link outage.', True, NAVY, 10)
doc.add_paragraph()

doc.add_heading('1. Executive plan', level=1)
doc.add_paragraph('Polar Twin already demonstrates selectable Maitri and Bharati profiles, storm and reset scenarios, a stock-out versus shipment outlook, and browser-side offline buffering. The next build should center on a complete, persistent response workflow. The API and SQLite database should own operational state; the web app should make that state clear to Station Leaders, Logistics, and Management. A schematic 3D station view, alert logic, health score, reports, and demo fallback then make the core workflow understandable and presentable.')
table(doc, ['Planning assumption', 'Plan'], [
    ('Delivery horizon', 'About 6–8 weeks for a small student team; adjust to team availability.'),
    ('Implementation stack', 'React + Vite web app; FastAPI API; SQLite for prototype persistence; Python simulator.'),
    ('Data posture', 'Synthetic or public-derived demo data only; provenance stays visible on every view and export.'),
    ('Primary demo story', 'Storm raises demand → fuel risk appears → delayed arrival creates a gap → request is approved and dispatched → ETA changes → delivery closes the incident.'),
])

doc.add_heading('2. Build sequence and milestones', level=1)
table(doc, ['Phase', 'Timing', 'Main deliverable', 'Exit check'], [
    ('0. Baseline and contracts', 'Days 1–3', 'Inventory existing code; agree API schemas, status names, provenance, and demo data.', 'Web and API run locally from a clean setup; contract examples reviewed.'),
    ('1. Workflow foundation', 'Week 1–2', 'SQLite-backed requisition, approval, dispatch, and delivery state machine.', 'Every transition is validated, persisted, timestamped, and visible after restart.'),
    ('2. Scenarios and stock outlook', 'Week 2–3', 'Generator fault during storm and delayed shipment scenarios connected to workflow ETA.', 'Repeatable scenario outputs; stock-out and ETA comparison updates from dispatch.'),
    ('3. Alerts and health', 'Week 3–4', 'Normal/Warning/Critical alerts with hysteresis and duration; health score breakdown.', 'No alert flicker in a scripted boundary sequence; score changes have clear reasons.'),
    ('4. Role views and station twin', 'Week 4–5', 'Role switcher and interactive schematic 3D/isometric station view.', 'Each role sees relevant actions; zone status maps to current scenario and alert.'),
    ('5. Reports and offline demo', 'Week 5–6', 'CSV/PDF export, seeded fallback mode, and prepared screenshots/recording.', 'Demo can complete without API connectivity using labeled seeded data.'),
    ('6. Hardening and handoff', 'Week 6–8', 'Pytest coverage, repeatability check, one-command start, CI workflow, README and license.', 'Clean-machine rehearsal succeeds; build and checks pass in CI.'),
])

doc.add_heading('3. Core requisition workflow', level=1)
doc.add_paragraph('Treat the workflow as a server-owned state machine. Each action should be an explicit API command, check the current state and actor role, write an audit event, and return the updated requisition. UI controls should reflect allowed transitions rather than allowing arbitrary status edits.')
table(doc, ['State', 'Owner', 'Action and required data', 'Next state'], [
    ('Draft', 'Station Leader', 'Select station and item; enter requested quantity, need-by date, and reason linked to alert/scenario.', 'Submitted'),
    ('Submitted', 'Logistics', 'Validate stock level, on-hand reserve, stock-out estimate, shipment availability, and requested quantity.', 'In review'),
    ('In review', 'Management', 'Approve, reject, or request changes; record approver, timestamp, and decision note.', 'Approved / Rejected / Changes requested'),
    ('Approved', 'Logistics', 'Create shipment; enter carrier/reference, dispatch time, expected arrival, and destination.', 'Dispatched'),
    ('Dispatched', 'Station Leader', 'Track ETA; show revised stock-out gap/buffer; record receipt quantity and delivery time.', 'Delivered'),
    ('Delivered', 'System + Station Leader', 'Reconcile inventory; close linked incident; retain full event history.', 'Closed'),
])
doc.add_heading('API and data design', level=2)
for text in [
    'Use a requisition record for current state and a separate append-only event record for every transition, comment, and ETA change.',
    'Suggested fields: requisition_id, station_id, item_id, quantity, unit, reason, linked_alert_id, status, requested_by, approved_by, shipment_reference, dispatched_at, eta_at, delivered_at, received_quantity, provenance, created_at, updated_at.',
    'Validate transitions on the API: for example, only an approved request may be dispatched, and only a dispatched shipment may be marked delivered.',
    'Dispatch must update the shipment ETA source used by the fuel stock-out comparison; retain the prior ETA in the audit history.',
    'Add role checks as lightweight demo rules now; replace with real authentication only after the core workflow is stable.',
]: bullet(doc, text)

doc.add_heading('4. Scenario and decision logic', level=1)
doc.add_heading('What-if scenarios', level=2)
table(doc, ['Scenario', 'Input change', 'Expected chain of effects'], [
    ('Generator fault during storm', 'Reduce generator availability while storm severity remains elevated.', 'Power margin falls → backup load increases → heating/fuel demand rises → alert and health score update → requisition can be raised.'),
    ('Delayed shipment', 'Move expected arrival later than the current shipment ETA.', 'Arrival buffer shrinks or shortfall appears → Logistics sees revised risk → approval/expedite decision is recorded.'),
])
doc.add_heading('Alert levels and hysteresis', level=2)
doc.add_paragraph('Define configurable demonstration thresholds for each key signal. Promote an alert only after its condition persists for a minimum duration. Clear or downgrade only after the signal remains below a lower recovery threshold for a separate duration. Store the rule version, observed inputs, duration, and explanation with each alert event. Label all thresholds illustrative until validated by domain experts.')
table(doc, ['Level', 'Purpose', 'Example behavior'], [
    ('Normal', 'Conditions within the configured operating band.', 'No active issue; retain history and show provenance.'),
    ('Warning', 'Sustained risk that needs review or planning.', 'Heating demand elevated or shipment buffer narrowing; create explainable action prompt.'),
    ('Critical', 'Sustained threshold breach with immediate operational attention in the demo.', 'Projected fuel stock-out before arrival or major loss of available power; highlight station zones and escalation path.'),
])
doc.add_heading('Station health score', level=2)
doc.add_paragraph('Show one 0–100 score with four visible components: energy, infrastructure, inventory, and environment. Normalize each component on an explicitly documented demo scale, combine with visible weights, and record the inputs and reason when the score changes. Avoid implying the score is a validated safety measure.')

doc.add_heading('5. User experience and station view', level=1)
table(doc, ['Role', 'Primary questions', 'Main actions'], [
    ('Station Leader', 'What is changing at the station? What needs a request or confirmation?', 'Review zone status; raise requisition; confirm delivery; view alert explanation.'),
    ('Logistics', 'Will inventory last until arrival? What must be dispatched or expedited?', 'Review requests; compare stock-out and ETA; set dispatch details; revise ETA.'),
    ('Management', 'What is the operational risk and what decision is waiting?', 'Review health and alert summary; approve, reject, or request changes.'),
])
doc.add_paragraph('Build the requested 3D station view as a lightweight isometric SVG schematic rather than a surveyed or navigable engineering model. Show named zones for habitat/heating, power generation, fuel storage, communications, and logistics. Color zones from live demo status, provide a legend and text labels, and offer a list-based equivalent for accessibility. Keep provenance and “schematic” labeling visible beside the view.')

doc.add_heading('6. Provenance, reports, and offline fallback', level=1)
for text in [
    'Make provenance required for every telemetry reading, scenario output, inventory estimate, requisition, shipment ETA, alert, score component, and export row.',
    'Use consistent values such as simulated, public-derived, manual, and sensor; display the source badge and last-updated time wherever a value appears.',
    'CSV should include units, station, time, provenance, scenario, and any linked workflow ID. PDF should summarize conditions, alert rationale, health components, inventory outlook, workflow history, and assumptions.',
    'Seed fallback mode with a normal baseline, storm, generator fault, and delayed shipment example. Make it available without a running API and label it clearly as seeded synthetic data.',
    'Save screenshots and a short screen recording of the complete story. Rehearse the live flow and fallback flow before presenting.',
]: bullet(doc, text)

doc.add_heading('7. Quality gates and delivery setup', level=1)
table(doc, ['Gate', 'Required evidence'], [
    ('API behavior', 'Pytest covers valid/invalid transitions, role permissions, persistence, ETA updates, and provenance validation.'),
    ('Simulator repeatability', 'Fixed seed produces equivalent outputs for the same profile and scenario.'),
    ('Frontend build', 'Production build succeeds; workflow and role views operate against API responses.'),
    ('Local run', 'One command starts the app and API, or a documented script handles prerequisites and startup.'),
    ('Continuous integration', 'GitHub Actions installs dependencies, runs API checks and simulator determinism check, and builds the web app.'),
    ('Repository documentation', 'README describes actual SQLite use, run steps, demo limits, provenance, and license choice.'),
])

doc.add_heading('8. Team ownership and working rhythm', level=1)
table(doc, ['Workstream', 'Suggested lead', 'Shared handoff'], [
    ('Web workflow, role views, 3D schematic, demo script', 'Mukesh / frontend lead', 'Agree on response schema with API lead; provide reviewable screens at each milestone.'),
    ('FastAPI, SQLite schema, state machine, exports', 'Tejas / API lead', 'Publish endpoint examples and transition rules before frontend integration.'),
    ('Simulator, scenario inputs, alert and score calculations', 'Anish / simulation lead', 'Document deterministic inputs, assumptions, thresholds, and explanations.'),
    ('Integration, release, README, CI', 'All three', 'Short integration check each day; rehearse one end-to-end path at each phase exit.'),
])
doc.add_paragraph('Keep a small shared board with To do, Doing, and Done. At the start of each work session, name one reviewable outcome. Share API contract changes immediately. At each phase exit, run the full demo flow together and move non-blocking enhancements to a later list.')

doc.add_heading('9. Demo run of show', level=1)
for title, detail in [
    ('Start at normal conditions', 'Select Maitri or Bharati and point out the synthetic provenance label and baseline health components.'),
    ('Trigger the storm', 'Show weather severity, heating load, fuel reserve, and alert explanation change together.'),
    ('Add a generator fault', 'Use the what-if to show reduced power margin, changed zone status, and the resulting health/alert reason.'),
    ('Raise a requisition', 'Station Leader submits a fuel request with a reason linked to the scenario.'),
    ('Review and approve', 'Logistics checks the projected stock-out against shipment ETA; Management records the decision.'),
    ('Dispatch and delay', 'Logistics records dispatch and ETA; delayed shipment scenario demonstrates a shrinking buffer or shortfall.'),
    ('Confirm delivery', 'Station Leader confirms receipt; inventory outlook updates and the workflow closes.'),
    ('Show resilience and evidence', 'Demonstrate offline buffering or seeded fallback, then export a provenance-labeled report.'),
]: numbered(doc, title, detail)

doc.add_heading('10. Risks and scope controls', level=1)
table(doc, ['Risk', 'Control'], [
    ('Prototype values are mistaken for live operational data.', 'Persistent synthetic labels, provenance fields, visible assumptions, and no live-station claims.'),
    ('Workflow becomes a set of disconnected UI buttons.', 'API owns state transitions; persist every action and audit event in SQLite.'),
    ('Too many features dilute the demo.', 'Finish requisition → dispatch → delivery first; add scenario coupling and explainability before polish.'),
    ('Thresholds or health score appear authoritative.', 'Mark as illustrative; expose formulas and require expert validation before operational use.'),
    ('Connectivity failure blocks presentation.', 'Provide seeded mode, screenshots, and screen recording as a rehearsed fallback.'),
])

doc.add_heading('Definition of a successful MVP demo', level=1)
doc.add_paragraph('A reviewer can run a repeatable scenario, understand the alert and health score, complete the requisition-to-delivery flow, and see dispatch update the stock-out outlook. Role views, the station schematic, provenance-labeled reporting, and offline fallback complete the demo.')

doc.core_properties.title = 'Polar Twin Project Building Plan'
doc.core_properties.subject = 'Phased MVP delivery plan for the Antarctic station operations prototype'
doc.core_properties.author = 'Polar Twin Team'
doc.save(OUT)
print(OUT)
